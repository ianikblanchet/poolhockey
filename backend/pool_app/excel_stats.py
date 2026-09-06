from io import BytesIO, StringIO
import csv
import re
import unicodedata

from openpyxl import load_workbook
from sqlalchemy.orm import Session

try:
    from .models import Player, PlayerSeasonStat
except ImportError:
    from models import Player, PlayerSeasonStat

DEFAULT_ACTIVE_SEASON = '2026-27'


def _excel_number(value, default=0):
    if value is None or value == '--':
        return default
    if isinstance(value, str):
        value = value.replace('\u202f', '').replace(' ', '').replace(',', '.')
    try:
        return float(value) if isinstance(value, str) and '.' in value else int(value)
    except (TypeError, ValueError):
        return default


def _excel_name(value):
    normalized = unicodedata.normalize('NFKD', str(value or '').casefold())
    without_accents = ''.join(char for char in normalized if not unicodedata.combining(char))
    return ' '.join(without_accents.strip().split())


def _normalize_season(season):
    season = (season or DEFAULT_ACTIVE_SEASON).strip().replace('/', '-')
    if len(season) == 4 and season.isdigit():
        year = int(season)
        return f'{year}-{str(year + 1)[2:]}'
    if len(season) in (5, 7) and '-' in season:
        return season
    return DEFAULT_ACTIVE_SEASON


def _season_from_filename(filename):
    match = re.search(r'(20\d{2})[-_]?(20\d{2})', filename or '')
    if match and int(match.group(2)) == int(match.group(1)) + 1:
        return f"{match.group(1)}-{match.group(2)[-2:]}"
    short_match = re.search(r'(?<!\d)(\d{2})[-_](\d{2})(?!\d)', filename or '')
    if short_match and int(short_match.group(2)) == int(short_match.group(1)) + 1:
        start_year = 2000 + int(short_match.group(1))
        return f'{start_year}-{short_match.group(2)}'
    return None


def _read_excel_rows(content):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    structured_rows = []
    embedded_csv_rows = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        if len(rows[0]) == 1:
            embedded_csv_rows.extend(
                list(csv.reader([str(row[0])]))[0]
                for row in rows
                if row and row[0] not in (None, '')
            )
        else:
            structured_rows.extend(rows)

    if embedded_csv_rows:
        header_index = next(
            (index for index, row in enumerate(embedded_csv_rows)
             if 'Player' in row and ('Season' in row or 'PTS' in row)),
            None,
        )
        if header_index is not None:
            return embedded_csv_rows[header_index:]
    return structured_rows


def import_active_stats_excel(db: Session, content: bytes, season_label=DEFAULT_ACTIVE_SEASON, filename=''):
    filename_season = _season_from_filename(filename)
    if filename_season:
        season_label = filename_season
    season_key = _normalize_season(season_label)
    start_year = int(season_key.split('-')[0])
    season_id = start_year * 10000 + start_year + 1
    players = {_excel_name(player.name): player for player in db.query(Player).all()}
    imported = 0
    created = 0
    unmatched = []
    wrong_season = 0
    found_header = False

    rows = _read_excel_rows(content)
    header = rows[0] if rows else None
    if header and 'PTS' in header and 'P' not in header:
        header = ['P' if value == 'PTS' else value for value in header]
        rows[0] = header
    if header and 'Player' in header:
        found_header = True
        columns = {name: index for index, name in enumerate(header)}
        save_percentage_column = 'Sv%' if 'Sv%' in columns else 'SV%' if 'SV%' in columns else None
        goalie = save_percentage_column is not None
        shots_column = columns.get('S', columns.get('SOG'))
        for row in rows[1:]:
            if not row:
                continue
            if 'Season' in columns and row[columns['Season']] not in (season_id, str(season_id)):
                wrong_season += 1
                continue
            player = players.get(_excel_name(row[columns['Player']]))
            if player is None:
                unmatched.append(str(row[columns['Player']]))
                continue

            stat = db.query(PlayerSeasonStat).filter(
                PlayerSeasonStat.player_id == player.id,
                PlayerSeasonStat.season == season_key,
            ).first()
            if stat is None:
                stat = PlayerSeasonStat(player_id=player.id, season=season_key)
                db.add(stat)
                created += 1

            stat.team = row[columns['Team']]
            stat.position = 'G' if goalie else row[columns['Pos']]
            stat.games_played = int(_excel_number(row[columns['GP']]))
            stat.wins = int(_excel_number(row[columns['W']])) if goalie and 'W' in columns else 0
            stat.goals = int(_excel_number(row[columns['G']]))
            stat.assists = int(_excel_number(row[columns['A']]))
            stat.points = int(_excel_number(row[columns['P']]))
            stat.plus_minus = 0.0 if goalie else float(_excel_number(row[columns['+/-']]))
            stat.shots = 0 if goalie or shots_column is None else int(_excel_number(row[shots_column]))
            stat.hits = 0
            stat.blocks = 0
            stat.pim = int(_excel_number(row[columns['PIM']]))
            if goalie and row[columns[save_percentage_column]] not in (None, '--'):
                stat.save_percentage = float(row[columns[save_percentage_column]])
            stat.source = 'nhl_excel'
            imported += 1

    if imported == 0:
        db.rollback()
        if not found_header:
            message = 'Le fichier Excel est vide ou ne contient pas les colonnes Player et Season.'
        else:
            message = f'Aucune ligne {season_key} trouvée dans le fichier NHL.'
        return {
            'status': 'warning',
            'message': message,
            'season': season_key,
            'rows_imported': 0,
            'rows_created': 0,
            'unmatched': unmatched,
            'wrong_season': wrong_season,
        }

    db.commit()
    return {
        'status': 'success',
        'message': f'{imported} statistiques NHL importées pour {season_key}.',
        'season': season_key,
        'rows_imported': imported,
        'rows_created': created,
        'unmatched': unmatched,
        'wrong_season': wrong_season,
    }


def _import_goalie_rows(db: Session, rows, season_key):
    players = {_excel_name(player.name): player for player in db.query(Player).all()}
    imported = 0
    created = 0
    unmatched = []

    if not rows:
        return {
            'status': 'warning',
            'message': 'Le fichier est vide ou ne contient pas d’en-tête.',
            'season': season_key,
            'rows_imported': 0,
            'rows_created': 0,
            'unmatched': [],
        }

    header_names = {str(name).strip() for name in (rows[0].keys() if rows else [])}
    player_column = 'Player' if 'Player' in header_names else None
    team_column = 'Team' if 'Team' in header_names else 'Tm' if 'Tm' in header_names else None
    required_columns = {'GP', 'W', 'SV%', 'G', 'A', 'PTS'}
    if player_column is None or team_column is None or not required_columns.issubset(header_names):
        return {
            'status': 'warning',
            'message': 'Le CSV doit contenir les colonnes Player, Tm ou Team, GP, W, SV%, G, A et PTS.',
            'season': season_key,
            'rows_imported': 0,
            'rows_created': 0,
            'unmatched': [],
        }

    for row in rows:
        player_name = (row.get(player_column) or '').strip()
        if not player_name or player_name == 'Player':
            continue
        player = players.get(_excel_name(player_name))
        if player is None:
            unmatched.append(player_name)
            continue

        stat = db.query(PlayerSeasonStat).filter(
            PlayerSeasonStat.player_id == player.id,
            PlayerSeasonStat.season == season_key,
        ).first()
        if stat is None:
            stat = PlayerSeasonStat(player_id=player.id, season=season_key)
            db.add(stat)
            created += 1

        stat.team = row.get(team_column)
        stat.position = 'G'
        stat.games_played = int(_excel_number(row.get('GP')))
        stat.wins = int(_excel_number(row.get('W')))
        stat.goals = int(_excel_number(row.get('G')))
        stat.assists = int(_excel_number(row.get('A')))
        stat.points = int(_excel_number(row.get('PTS')))
        stat.plus_minus = 0.0
        stat.shots = 0
        stat.hits = 0
        stat.blocks = 0
        stat.pim = int(_excel_number(row.get('PIM')))
        stat.save_percentage = float(_excel_number(row.get('SV%'))) if row.get('SV%') not in (None, '', '--') else None
        stat.source = 'hockey_reference_csv'
        imported += 1

    if imported == 0:
        db.rollback()
        return {
            'status': 'warning',
            'message': 'Aucune statistique de gardien correspondante n’a été importée.',
            'season': season_key,
            'rows_imported': 0,
            'rows_created': 0,
            'unmatched': unmatched,
        }

    db.commit()
    return {
        'status': 'success',
        'message': f'{imported} statistiques de gardiens importées pour {season_key}.',
        'season': season_key,
        'rows_imported': imported,
        'rows_created': created,
        'unmatched': unmatched,
    }


def import_goalie_stats_csv(db: Session, content: bytes, season_label=DEFAULT_ACTIVE_SEASON):
    season_key = _normalize_season(season_label)
    text = content.decode('utf-8-sig')
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.DictReader(StringIO(text), dialect=dialect))
    return _import_goalie_rows(db, rows, season_key)


def import_goalie_stats_excel(db: Session, content: bytes, season_label=DEFAULT_ACTIVE_SEASON):
    season_key = _normalize_season(season_label)
    raw_rows = _read_excel_rows(content)
    if not raw_rows:
        return _import_goalie_rows(db, [], season_key)

    header = [str(value).strip() if value is not None else '' for value in raw_rows[0]]
    rows = [dict(zip(header, row)) for row in raw_rows[1:] if row]
    return _import_goalie_rows(db, rows, season_key)
